"""Turn a DCF result into a console summary, an Excel workbook and heatmap charts."""
from __future__ import annotations

from pathlib import Path

import pandas as pd

from .data import CompanyData
from .model import Assumptions, DCFResult


def _money(x: float) -> str:
    for unit, div in (("T", 1e12), ("B", 1e9), ("M", 1e6)):
        if abs(x) >= div:
            return f"{x / div:,.2f}{unit}"
    return f"{x:,.0f}"


def _pct_table(table: pd.DataFrame) -> pd.DataFrame:
    """Format index/columns as percentages for display."""
    t = table.copy()
    t.index = [f"{v:.1%}" for v in t.index]
    t.columns = [f"{v:.1%}" for v in t.columns]
    return t


# ---------------------------------------------------------------------------
# Console
# ---------------------------------------------------------------------------

def print_summary(data: CompanyData, a: Assumptions, r: DCFResult,
                  sens_tg: pd.DataFrame, sens_g: pd.DataFrame) -> None:
    line = "=" * 64
    print(f"\n{line}\n  DCF VALUATION: {data.name} ({data.ticker})\n{line}")

    print("\nKey assumptions")
    print(f"  Starting unlevered FCF   {_money(r.base_fcf)} {data.currency}")
    print(f"  Initial growth           {r.initial_growth:.2%}  ({r.growth_source})")
    print(f"  Growth path              held {a.hold_years} yrs, then fades to terminal by yr "
          f"{a.projection_years}")
    print(f"  Terminal growth          {a.terminal_growth:.2%}")
    if r.wacc_breakdown:
        w = r.wacc_breakdown
        print(f"  Risk-free rate           {data.risk_free_rate:.2%}")
        beta_note = " (Blume-adjusted)" if a.adjust_beta else ""
        print(f"  Beta                     {data.beta:.2f} raw -> {w.beta_used:.2f} used{beta_note}")
        print(f"  Cost of equity (CAPM)    {w.cost_of_equity:.2%}")
        print(f"  After-tax cost of debt   {w.after_tax_cost_of_debt:.2%}")
        print(f"  Weights (E / D)          {w.weight_equity:.0%} / {w.weight_debt:.0%}")
    print(f"  WACC                     {r.wacc:.2%}")

    proj = r.projections.copy()
    proj["Growth"] = proj["Growth"].map("{:.2%}".format)
    proj["FCF"] = proj["FCF"].map(_money)
    proj["Discount factor"] = proj["Discount factor"].map("{:.4f}".format)
    proj["PV of FCF"] = proj["PV of FCF"].map(_money)
    print("\nProjected free cash flow")
    print(proj.to_string(index=False))

    print("\nValuation")
    print(f"  PV of projected FCF      {_money(r.projections['PV of FCF'].sum())}")
    print(f"  PV of terminal value     {_money(r.pv_terminal_value)}"
          f"  ({r.terminal_share_of_ev:.0%} of EV)")
    print(f"  Enterprise value         {_money(r.enterprise_value)}")
    print(f"  - Debt                   {_money(data.total_debt)}")
    print(f"  + Cash                   {_money(data.cash)}")
    print(f"  Equity value             {_money(r.equity_value)}")
    print(f"  Shares outstanding       {_money(data.shares_outstanding)}")

    verdict = "UNDERVALUED" if r.upside > 0 else "OVERVALUED"
    print(f"\n  Intrinsic value / share  {r.value_per_share:,.2f} {data.currency}")
    print(f"  Current price            {data.price:,.2f} {data.currency}")
    print(f"  Upside / (downside)      {r.upside:+.1%}  -> model says {verdict}")

    print("\nReverse DCF: what is the market pricing in?")
    if r.implied_growth is not None:
        print(f"  To justify {data.price:,.2f}, FCF must grow ~{r.implied_growth:.1%} a year for "
              f"{a.hold_years} yrs, then fade to {a.terminal_growth:.1%} by yr {a.projection_years}.")
        print(f"  Your assumption: {r.initial_growth:.1%}. Ask yourself which is more realistic.")
    else:
        print("  No growth rate between -20% and 100% justifies the current price "
              "(check the WACC and starting FCF).")
    print(f"  Market EV / FCF today    {r.market_ev_to_fcf:,.1f}x")
    print(f"  Model's exit EV / FCF    {r.terminal_ev_to_fcf:,.1f}x  (multiple the terminal value implies)")

    print("\nSensitivity: value per share (rows = WACC, columns = terminal growth)")
    print(_pct_table(sens_tg).round(2).to_string())
    print("\nSensitivity: value per share (rows = WACC, columns = initial growth)")
    print(_pct_table(sens_g).round(2).to_string())

    if r.terminal_share_of_ev > 0.8:
        print("\n  Note: terminal value is >80% of EV, so the result is very "
              "sensitive to long-run assumptions.")
    print()


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def write_excel(path: Path, data: CompanyData, a: Assumptions, r: DCFResult,
                sens_tg: pd.DataFrame, sens_g: pd.DataFrame) -> Path:
    from openpyxl.styles import Font, PatternFill

    summary = pd.DataFrame([
        ("Company", data.name),
        ("Ticker", data.ticker),
        ("Currency", data.currency),
        ("Current price", data.price),
        ("Intrinsic value per share", r.value_per_share),
        ("Upside / (downside)", r.upside),
        ("Enterprise value", r.enterprise_value),
        ("Equity value", r.equity_value),
        ("PV of terminal value % of EV", r.terminal_share_of_ev),
        ("Implied growth (reverse DCF)", r.implied_growth if r.implied_growth is not None else "n/a"),
        ("Market EV / FCF", r.market_ev_to_fcf),
        ("Model exit EV / FCF", r.terminal_ev_to_fcf),
        ("", ""),
        ("Starting unlevered FCF", r.base_fcf),
        ("Initial growth", r.initial_growth),
        ("Growth source", r.growth_source),
        ("Years growth is held", a.hold_years),
        ("Terminal growth", a.terminal_growth),
        ("Projection years", a.projection_years),
        ("WACC", r.wacc),
        ("Risk-free rate", data.risk_free_rate),
        ("Beta (raw)", data.beta),
        ("Beta (used)", r.wacc_breakdown.beta_used if r.wacc_breakdown else "n/a"),
        ("Equity risk premium", a.equity_risk_premium),
        ("Tax rate", data.tax_rate),
        ("Total debt", data.total_debt),
        ("Cash", data.cash),
        ("Shares outstanding", data.shares_outstanding),
    ], columns=["Item", "Value"])

    history = pd.DataFrame({
        "Revenue": data.revenue_history,
        "Free cash flow": data.fcf_history,
    })
    history.index = [getattr(i, "year", i) for i in history.index]
    history.index.name = "Fiscal year"

    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        r.projections.to_excel(xw, sheet_name="Projections", index=False)
        history.to_excel(xw, sheet_name="History")
        sens_tg.to_excel(xw, sheet_name="Sens WACC x TG")
        sens_g.to_excel(xw, sheet_name="Sens WACC x Growth")

        green = PatternFill("solid", fgColor="C6EFCE")
        red = PatternFill("solid", fgColor="FFC7CE")
        bold = Font(bold=True)

        for name in ("Sens WACC x TG", "Sens WACC x Growth"):
            ws = xw.sheets[name]
            for row in ws.iter_rows(min_row=2, min_col=2):
                for cell in row:
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0.00"
                        cell.fill = green if cell.value >= data.price else red
            for cell in ws[1][1:]:
                cell.number_format = "0.0%"
                cell.font = bold
            for row in ws.iter_rows(min_row=2, max_col=1):
                row[0].number_format = "0.0%"
                row[0].font = bold
            ws.cell(row=ws.max_row + 2, column=1,
                    value=f"Green = above current price ({data.price:,.2f}); red = below")

        ws = xw.sheets["Projections"]
        for row in ws.iter_rows(min_row=2):
            row[1].number_format = "0.00%"
            row[2].number_format = "#,##0"
            row[3].number_format = "0.0000"
            row[4].number_format = "#,##0"

        ws = xw.sheets["Summary"]
        pct_rows = {"Upside / (downside)", "PV of terminal value % of EV", "Initial growth",
                    "Implied growth (reverse DCF)",
                    "Terminal growth", "WACC", "Risk-free rate", "Equity risk premium", "Tax rate"}
        for row in ws.iter_rows(min_row=2):
            label, val = row[0].value, row[1]
            if label in pct_rows:
                val.number_format = "0.00%"
            elif isinstance(val.value, float):
                val.number_format = "#,##0.00"
        ws.column_dimensions["A"].width = 32
        ws.column_dimensions["B"].width = 22
    return path


# ---------------------------------------------------------------------------
# Heatmap
# ---------------------------------------------------------------------------

def plot_heatmap(path: Path, table: pd.DataFrame, price: float, title: str,
                 col_label: str) -> Path:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.colors import TwoSlopeNorm

    values = table.to_numpy(dtype=float)
    finite = values[~pd.isna(values)]
    vmin = min(finite.min(), price * 0.99)
    vmax = max(finite.max(), price * 1.01)
    norm = TwoSlopeNorm(vmin=vmin, vcenter=price, vmax=vmax)

    fig, ax = plt.subplots(figsize=(8, 5.5))
    im = ax.imshow(values, cmap="RdYlGn", norm=norm, aspect="auto")

    ax.set_xticks(range(len(table.columns)), [f"{c:.1%}" for c in table.columns])
    ax.set_yticks(range(len(table.index)), [f"{i:.1%}" for i in table.index])
    ax.set_xlabel(col_label)
    ax.set_ylabel("WACC")
    ax.set_title(f"{title}\n(current price {price:,.2f}; green = above, red = below)")

    mid_r, mid_c = len(table.index) // 2, len(table.columns) // 2
    for i in range(values.shape[0]):
        for j in range(values.shape[1]):
            v = values[i, j]
            text = "n/a" if pd.isna(v) else f"{v:,.0f}"
            weight = "bold" if (i, j) == (mid_r, mid_c) else "normal"
            ax.text(j, i, text, ha="center", va="center", fontsize=9, weight=weight)

    fig.colorbar(im, ax=ax, label="Value per share")
    fig.tight_layout()
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return path
